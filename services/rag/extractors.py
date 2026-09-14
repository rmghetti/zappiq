"""
Conversores de arquivo para texto, com as recusas que o cliente consegue ler.

Por que existe um modulo so para isto:
  1. A tela do Treinar IA sempre prometeu PDF, Word, Excel, texto, Markdown e
     CSV. O servico so lia PDF e text/*. Word e Excel morriam com 415 aqui,
     viravam "Internal Server Error" na API e o cliente nao tinha o que fazer
     com a mensagem (achados A003 e A004).
  2. O leitor de PDF era o pymupdf, AGPL-3.0. Num servico de rede
     multi-inquilino a AGPL pede oferecer o codigo-fonte a quem usa o servico
     (achado A235). Aqui o PDF passa a ser lido pelo pypdf, BSD-3-Clause.
  3. Pagina de site entrava crua: o site do CMJ rendeu 1 trecho de 29
     caracteres e ainda assim apareceu como indexado, e perfil de rede social
     virou menu de navegacao no vetor (achado A012).

Regra de ouro deste modulo: toda recusa sai com status HTTP correto e com uma
frase em portugues que diz o que fazer em seguida. Nada de 500 generico.
"""

import io
import re
from urllib.parse import urlparse

from fastapi import HTTPException

# ─────────────────────────────────────────────────────────────────────────────
# Mensagens (unica fonte: a API repassa estas frases ao cliente)
# ─────────────────────────────────────────────────────────────────────────────

MENSAGEM_TIPO_NAO_ACEITO = (
    "Tipo de arquivo não aceito: envie PDF, Word (.docx), Excel (.xlsx), "
    "texto, Markdown ou CSV."
)
MENSAGEM_ARQUIVO_ILEGIVEL = (
    "Não consegui ler este arquivo. Salve como PDF ou texto e envie de novo."
)
MENSAGEM_PDF_SEM_TEXTO = (
    "Este PDF parece digitalizado, sem texto selecionável. "
    "Envie um PDF com texto ou cole o conteúdo."
)
MENSAGEM_SEM_TEXTO = (
    "Não encontrei texto neste arquivo. Confira o conteúdo e envie de novo."
)
MENSAGEM_PAGINA_CURTA = (
    "A página não tem texto legível o suficiente. Cole o conteúdo como texto."
)
MENSAGEM_REDE_SOCIAL = (
    "Redes sociais não podem ser lidas automaticamente. "
    "Cole o texto do perfil ou da publicação."
)


def mensagem_arquivo_grande(tamanho_mb: float, limite_mb: int) -> str:
    """O limite entra na frase: sem o numero o cliente nao sabe quanto cortar."""
    return (
        f"Arquivo de {tamanho_mb:.1f} MB acima do limite de {limite_mb} MB. "
        "Divida o arquivo ou envie uma versão menor."
    )


# Menos que isto nao e conteudo de pagina, e menu de navegacao.
MINIMO_CARACTERES_PAGINA = 200

# Dominios que nao entregam conteudo a um leitor sem sessao: o que volta e
# titulo com entidade HTML e menu. Comparado por dominio registravel, nunca por
# substring (senao "meufacebook.com.br" cairia junto).
DOMINIOS_REDE_SOCIAL = frozenset(
    {
        "instagram.com",
        "facebook.com",
        "fb.com",
        "tiktok.com",
        "x.com",
        "twitter.com",
        "linkedin.com",
        "youtube.com",
        "youtu.be",
    }
)


# ─────────────────────────────────────────────────────────────────────────────
# Deteccao de formato
# ─────────────────────────────────────────────────────────────────────────────

_EXTENSOES = {
    ".pdf": "pdf",
    ".docx": "docx",
    ".xlsx": "xlsx",
    ".txt": "texto",
    ".md": "texto",
    ".markdown": "texto",
    ".csv": "texto",
    ".html": "html",
    ".htm": "html",
}

_MIMES = {
    "application/pdf": "pdf",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": "docx",
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet": "xlsx",
    "text/html": "html",
    "application/xhtml+xml": "html",
}


def detectar_formato(content_type: str | None, filename: str) -> str | None:
    """
    Descobre o formato pela extensao e, se ela nao disser nada, pelo mime.

    A extensao vem primeiro de proposito: o navegador manda
    application/vnd.ms-excel para .csv e application/octet-stream para quase
    tudo que ele nao reconhece. A extensao e o que o cliente ve e escolheu.
    """
    nome = (filename or "").strip().lower()
    for extensao, formato in _EXTENSOES.items():
        if nome.endswith(extensao):
            return formato

    mime = (content_type or "").split(";")[0].strip().lower()
    if mime in _MIMES:
        return _MIMES[mime]
    if mime.startswith("text/"):
        return "texto"
    return None


# ─────────────────────────────────────────────────────────────────────────────
# Conversores
# ─────────────────────────────────────────────────────────────────────────────


def _ilegivel() -> HTTPException:
    return HTTPException(status_code=422, detail=MENSAGEM_ARQUIVO_ILEGIVEL)


def _extrair_pdf(data: bytes) -> str:
    """PDF com pypdf (BSD-3-Clause). Quebra de pagina vira linha em branco."""
    from pypdf import PdfReader

    try:
        leitor = PdfReader(io.BytesIO(data))
        paginas = [(pagina.extract_text() or "") for pagina in leitor.pages]
    except Exception as exc:
        raise _ilegivel() from exc

    texto = "\n\n".join(p.strip() for p in paginas if p.strip()).strip()
    if not texto:
        raise HTTPException(status_code=422, detail=MENSAGEM_PDF_SEM_TEXTO)
    return texto


def _extrair_docx(data: bytes) -> str:
    """Word com mammoth (BSD-2-Clause). Texto puro: estilo nao ajuda a busca."""
    import mammoth

    try:
        resultado = mammoth.extract_raw_text(io.BytesIO(data))
    except Exception as exc:
        raise _ilegivel() from exc
    return (resultado.value or "").strip()


def _extrair_xlsx(data: bytes) -> str:
    """
    Excel com openpyxl (MIT). Uma linha de texto por linha da planilha, com o
    nome da aba como secao. Celula vazia sai fora: "Curso | | 1200" nao ajuda
    ninguem, e o cabecalho de trecho usa o titulo da aba para se situar.
    """
    import openpyxl

    try:
        planilha = openpyxl.load_workbook(
            io.BytesIO(data), read_only=True, data_only=True
        )
    except Exception as exc:
        raise _ilegivel() from exc

    linhas: list[str] = []
    try:
        for aba in planilha.worksheets:
            linhas.append(f"## Planilha: {aba.title}")
            for celulas in aba.iter_rows(values_only=True):
                valores = [
                    str(celula).strip()
                    for celula in celulas
                    if celula is not None and str(celula).strip()
                ]
                if valores:
                    linhas.append(" | ".join(valores))
    except Exception as exc:
        raise _ilegivel() from exc
    finally:
        planilha.close()

    return "\n".join(linhas).strip()


def _extrair_texto_puro(data: bytes) -> str:
    try:
        return data.decode("utf-8", errors="replace").strip()
    except Exception as exc:
        raise _ilegivel() from exc


def _extrair_html(data: bytes) -> str:
    """
    Conteudo principal da pagina com Readability (Apache-2.0), o mesmo algoritmo
    do modo leitura do navegador: sai menu, rodape, script e banner de cookie.
    """
    from lxml import html as lxml_html
    from readability import Document

    bruto = data.decode("utf-8", errors="replace")
    try:
        documento = Document(bruto)
        principal = documento.summary(html_partial=True)
        arvore = lxml_html.fromstring(principal)
        texto = arvore.text_content()
    except Exception:
        # Readability desiste de pagina sem corpo reconhecivel. Nao e erro de
        # leitura do arquivo: o portao de conteudo minimo abaixo decide.
        texto = ""

    if not texto.strip():
        try:
            texto = lxml_html.fromstring(bruto).text_content()
        except Exception:
            texto = ""

    limpo = re.sub(r"[ \t\xa0]+", " ", texto)
    limpo = re.sub(r"\s*\n\s*", "\n", limpo).strip()

    if len(re.sub(r"\s", "", limpo)) < MINIMO_CARACTERES_PAGINA:
        raise HTTPException(status_code=422, detail=MENSAGEM_PAGINA_CURTA)
    return limpo


# ─────────────────────────────────────────────────────────────────────────────
# Portao de rede social
# ─────────────────────────────────────────────────────────────────────────────


def verificar_rede_social(url: str | None) -> None:
    """Levanta 422 se a URL for de rede social. Sem URL, nao ha o que checar."""
    if not url:
        return
    try:
        host = (urlparse(url).hostname or "").lower()
    except ValueError:
        return
    if not host:
        return

    partes = host.split(".")
    # Compara o dominio registravel e os sufixos: "br.linkedin.com" e
    # "m.facebook.com" entram; "meufacebook.com.br" nao.
    for corte in range(len(partes) - 1):
        if ".".join(partes[corte:]) in DOMINIOS_REDE_SOCIAL:
            raise HTTPException(status_code=422, detail=MENSAGEM_REDE_SOCIAL)


# ─────────────────────────────────────────────────────────────────────────────
# Porta unica
# ─────────────────────────────────────────────────────────────────────────────


def extrair_texto(
    content_type: str | None,
    filename: str,
    data: bytes,
    source_url: str | None = None,
) -> str:
    """
    Converte o arquivo em texto ou levanta HTTPException com mensagem util.

    415 tipo nao aceito, 422 arquivo ilegivel, PDF sem texto, pagina curta ou
    rede social. Nunca 500 por formato de entrada.
    """
    verificar_rede_social(source_url)

    formato = detectar_formato(content_type, filename)
    if formato is None:
        raise HTTPException(status_code=415, detail=MENSAGEM_TIPO_NAO_ACEITO)

    if formato == "pdf":
        return _extrair_pdf(data)
    if formato == "docx":
        return _extrair_docx(data)
    if formato == "xlsx":
        return _extrair_xlsx(data)
    if formato == "html":
        return _extrair_html(data)
    return _extrair_texto_puro(data)
